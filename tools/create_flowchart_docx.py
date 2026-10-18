from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Cm


ROOT = Path(__file__).resolve().parents[1]
ARTIFACTS = ROOT / "artifacts"
PNG = ARTIFACTS / "questionnaire-system-flowchart.png"
OUT = ARTIFACTS / "questionnaire-system-flowchart.docx"


doc = Document()
section = doc.sections[0]
section.page_width = Cm(21.0)
section.page_height = Cm(29.7)
section.top_margin = Cm(1.4)
section.bottom_margin = Cm(1.4)
section.left_margin = Cm(1.7)
section.right_margin = Cm(1.7)

figure = doc.add_paragraph()
figure.alignment = WD_ALIGN_PARAGRAPH.CENTER
figure.add_run().add_picture(str(PNG), width=Cm(14.7))

doc.save(OUT)
print(OUT)
