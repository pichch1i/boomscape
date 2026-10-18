from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "artifacts" / "questionnaire-system-flowchart-editable.docx"


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_border(cell, color="2D5B88", size="18"):
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = tc_pr.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    for edge in ("top", "left", "bottom", "right"):
        tag = "w:" + edge
        element = borders.find(qn(tag))
        if element is None:
            element = OxmlElement(tag)
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), size)
        element.set(qn("w:color"), color)


def clear_cell_border(cell):
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = tc_pr.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = "w:" + edge
        element = borders.find(qn(tag))
        if element is None:
            element = OxmlElement(tag)
            borders.append(element)
        element.set(qn("w:val"), "nil")


def set_cell_margins(cell, top=80, start=110, bottom=80, end=110):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for margin, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{margin}"))
        if node is None:
            node = OxmlElement(f"w:{margin}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_run_font(run, size=12, bold=False, color="172033"):
    run.font.name = "Tahoma"
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = RGBColor.from_string(color)
    fonts = run._element.get_or_add_rPr().get_or_add_rFonts()
    for key in ("ascii", "hAnsi", "eastAsia", "cs"):
        fonts.set(qn(f"w:{key}"), "Tahoma")


def set_text(cell, text, size=12, bold=False, color="172033", align=WD_ALIGN_PARAGRAPH.CENTER):
    cell.text = ""
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    paragraph = cell.paragraphs[0]
    paragraph.alignment = align
    paragraph.paragraph_format.space_before = Pt(0)
    paragraph.paragraph_format.space_after = Pt(0)
    paragraph.paragraph_format.line_spacing = 1.0
    lines = text.split("\n")
    for index, line in enumerate(lines):
        run = paragraph.add_run(line)
        set_run_font(run, size=size, bold=bold, color=color)
        if index < len(lines) - 1:
            run.add_break()


def set_row_height(row, cm):
    tr_pr = row._tr.get_or_add_trPr()
    height = OxmlElement("w:trHeight")
    height.set(qn("w:val"), str(int(cm * 567)))
    height.set(qn("w:hRule"), "atLeast")
    tr_pr.append(height)


doc = Document()
section = doc.sections[0]
section.page_width = Cm(21.0)
section.page_height = Cm(29.7)
section.top_margin = Cm(0.6)
section.bottom_margin = Cm(0.6)
section.left_margin = Cm(1.4)
section.right_margin = Cm(1.4)

title = doc.add_paragraph()
title.alignment = WD_ALIGN_PARAGRAPH.CENTER
title.paragraph_format.space_after = Pt(2)
title_run = title.add_run("ผังงานลำดับการทำงานของระบบแบบสอบถาม")
set_run_font(title_run, size=14, bold=True, color="111827")

table = doc.add_table(rows=0, cols=3)
table.alignment = WD_TABLE_ALIGNMENT.CENTER
table.autofit = False
table.columns[0].width = Cm(3.1)
table.columns[1].width = Cm(11.9)
table.columns[2].width = Cm(3.1)


def add_node(text, kind="process", left_note="", right_note=""):
    row = table.add_row()
    set_row_height(row, 0.78 if "\n" not in text else 0.98)
    for idx, width in enumerate((3.1, 11.9, 3.1)):
        row.cells[idx].width = Cm(width)
        clear_cell_border(row.cells[idx])
        set_cell_margins(row.cells[idx], top=28, bottom=28)

    center = row.cells[1]
    if kind == "terminal":
        set_cell_shading(center, "275A82")
        set_cell_border(center, "275A82", "18")
        set_text(center, text, size=10.5, bold=True, color="FFFFFF")
    elif kind == "decision":
        set_cell_shading(center, "FFF5DF")
        set_cell_border(center, "B46B22", "18")
        set_text(center, "◆  " + text + "  ◆", size=9.5, bold=True)
    elif kind == "result":
        set_cell_shading(center, "F3EDFF")
        set_cell_border(center, "7057A3", "18")
        set_text(center, text, size=10)
    elif kind == "backend":
        set_cell_shading(center, "EAF7EF")
        set_cell_border(center, "2C7A52", "18")
        set_text(center, text, size=10)
    elif kind == "id":
        set_cell_shading(center, "EAF4FF")
        set_cell_border(center, "2D5B88", "18")
        set_text(center, text, size=10)
    else:
        set_cell_shading(center, "F7F9FC")
        set_cell_border(center, "2D5B88", "18")
        set_text(center, text, size=10)

    if left_note:
        set_text(row.cells[0], left_note, size=8.5, bold=True, color="B42318")
    if right_note:
        set_text(row.cells[2], right_note, size=8.5, bold=True, color="2C7A52")


def add_arrow(label=""):
    row = table.add_row()
    set_row_height(row, 0.25)
    for cell in row.cells:
        clear_cell_border(cell)
        set_cell_margins(cell, top=0, bottom=0)
    arrow = "↓" if not label else f"{label}  ↓"
    set_text(row.cells[1], arrow, size=9.5, bold=True, color="44546A")


add_node("เริ่มต้น", "terminal")
add_arrow()
add_node("แสดงหน้าเริ่มต้นและรายละเอียดแบบสอบถาม")
add_arrow()
add_node("กรอกข้อมูลผู้เล่น\nและให้ความยินยอมในการจัดเก็บข้อมูล")
add_arrow()
add_node("ข้อมูลครบถ้วนและให้ความยินยอมแล้วหรือไม่", "decision", "ไม่ใช่ ↩\nกลับไปกรอกข้อมูล", "ใช่")
add_arrow()
add_node("สร้างรหัสประจำการตอบแบบสอบถาม\n(submissionId)", "id")
add_arrow()
add_node("แสดงคำถามและบันทึกคำตอบข้อที่ 1–7")
add_arrow()
add_node("ตอบคำถามครบทั้ง 7 ข้อหรือไม่", "decision", "ไม่ใช่ ↩\nกลับไปตอบคำถาม", "ใช่")
add_arrow()
add_node("รวมคะแนนอารมณ์และเลือกผลลัพธ์")
add_arrow()
add_node("แสดงดอกไม้ ข้อความ และผลลัพธ์ของผู้เล่น", "result")
add_arrow()
add_node("บันทึกข้อมูลลงฐานข้อมูล Supabase", "backend")
add_arrow()
add_node("ส่งต่อข้อมูลสำหรับแสดงผลในหน้า Admin\nและเชื่อมต่อกับ TouchDesigner", "backend")
add_arrow()
add_node("สิ้นสุด", "terminal")

caption = doc.add_paragraph()
caption.alignment = WD_ALIGN_PARAGRAPH.CENTER
caption.paragraph_format.space_before = Pt(2)
caption_run = caption.add_run("ภาพที่ 3.X ผังงานลำดับการทำงานของระบบแบบสอบถาม")
set_run_font(caption_run, size=10.5, color="111827")

OUT.parent.mkdir(parents=True, exist_ok=True)
doc.save(OUT)
print(OUT)
